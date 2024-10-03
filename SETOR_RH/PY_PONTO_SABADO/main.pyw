import pandas as pd
import os
import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import datetime
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from time import sleep

# função para exibir uma janela de erro
def mostrar_erro(mensagem):
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("Erro", mensagem)
    root.destroy()

# função para selecionar os arquivos a serem analisados
def selecionar_arquivos():
    root = tk.Tk()
    root.withdraw()
    caminhos_arquivos = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx")])
    return caminhos_arquivos

# cria a janela principal
root = tk.Tk()
root.withdraw()  # oculta a janela principal

# variaveis da data
date = datetime.now()
mes = f'{date.month:02d}'
dia = f'{date.day:02d}'

# pede pro usuario selecionar as planilhas 
messagebox.showinfo("Selecionar Arquivo", "Selecione a planilha base dos colaboradores!")
base_colaboradores = selecionar_arquivos()
if not base_colaboradores:
    mostrar_erro(f"Nenhum arquivo selecionado. Encerrando programa!")
    exit()

messagebox.showinfo("Selecionar Arquivo", "Selecione todas planilhas de ponto!")
planilhas = selecionar_arquivos()
if not planilhas:
    mostrar_erro(f"Nenhum arquivo selecionado. Encerrando programa!")
    exit()

# variavel para armazenar os pontos dos colaboradores
presenca = {}

# analisa cada planilha de ponto
for i, planilha in enumerate(planilhas):
    try:
        # ler a planilha
        df = pd.read_excel(planilha)

        # verificar se a primeira coluna (index 0) existe e tentar extrair a data
        if df.shape[1] > 0:
            primeira_coluna = str(df.iloc[0, 0])
            data_match = re.search(r'(\d{2}/\d{2}/\d{4})', primeira_coluna)

            if data_match:
                data_str = data_match.group(1)
                data_formatada = datetime.strptime(data_str, '%d/%m/%Y').strftime('%d-%m-%Y')
            else:
                data_formatada = 'data_nao_encontrada'

            # verificar se a segunda coluna (index 1) existe
            if df.shape[1] > 1:
                for index, valor in enumerate(df.iloc[:, 1]):  # acessa a segunda coluna
                    if isinstance(valor, str) and not any(char.isdigit() for char in valor):
                        nome = valor.strip()

                        # verifica as linhas abaixo do nome do colaborador
                        if index + 2 < len(df):
                            linha_baixo_1 = df.iloc[index + 1, 1]
                            linha_baixo_2 = df.iloc[index + 2, 1]

                            if (pd.notna(linha_baixo_1) and linha_baixo_1 != '') and \
                               (pd.notna(linha_baixo_2) and linha_baixo_2 != ''):
                                if nome not in presenca:
                                    presenca[nome] = [0] * len(planilhas)

                                presenca[nome][i] = 1
                            else:
                                continue
            else:
                mostrar_erro(f"A planilha {planilha} não contém a coluna esperada.")
    except Exception as e:
        mostrar_erro(f"erro ao processar {planilha}: {e}")

# filtrar apenas colaboradores que tem pelo menos uma presença
presenca_filtrada = {k: v for k, v in presenca.items() if any(presenca[k])}

# criar um DataFrame a partir do dicionario filtrado
df_nomes = pd.DataFrame.from_dict(presenca_filtrada, orient='index', columns=[f'planilha_{i + 1}' for i in range(len(planilhas))]).reset_index()
df_nomes.columns = ['colaborador'] + [datetime.strptime(re.search(r"(\d{2}/\d{2}/\d{4})", str(pd.read_excel(planilhas[j]).iloc[0, 0])).group(1), "%d/%m/%Y").strftime("%d-%m-%Y") for j in range(len(planilhas))]

# adicionar a coluna total
df_nomes['total'] = df_nomes.iloc[:, 1:].sum(axis=1)

# ordenar o DataFrame pelo nome do colaborador
df_nomes.sort_values(by='colaborador', inplace=True)

# processar a planilha base para obter os cargos
try:
    # ler a planilha base com o cabeçalho na quarta linha
    base_df = pd.read_excel(base_colaboradores[0], header=3)
    
    # selecionar as colunas 'I' e 'J'
    col_colaborador = base_df.columns[8]
    col_cargo = base_df.columns[9]  

    # adicionar uma nova coluna para cargos
    cargo_dict = dict(zip(base_df[col_colaborador], base_df[col_cargo]))

    # verificando se todos os colaboradores encontrados têm um cargo
    df_nomes['cargo'] = df_nomes['colaborador'].map(cargo_dict).fillna('cargo_nao_encontrado')

    # reordenar as colunas para que 'cargo' fique ao lado de 'colaborador'
    df_nomes = df_nomes[['colaborador', 'cargo'] + [col for col in df_nomes.columns if col not in ['colaborador', 'cargo']]]
    
except Exception as e:
    mostrar_erro(f"erro ao processar a planilha base de colaboradores: {e}")
    

# salvar o DataFrame em uma nova planilha na pasta de downloads do usuario
usuario = os.getlogin()
pasta = f'C:/Users/{usuario}/Downloads/Colaboradores_Ponto_Sábados_{dia}_{mes}.xlsx'

# formatando a tabela para melhor visualização
with pd.ExcelWriter(pasta, engine='openpyxl') as writer:
    df_nomes.to_excel(writer, index=False, sheet_name='Colaboradores')
    workbook = writer.book
    worksheet = writer.sheets['Colaboradores']

    # aplicar formatação em listras
    for row in range(2, len(df_nomes) + 2):  # começar da linha 2, pois a linha 1 é o cabeçalho
        if row % 2 == 0:
            for col in range(1, len(df_nomes.columns) + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    # formatar cabeçalho
    for col in range(1, len(df_nomes.columns) + 1):
        header_cell = worksheet.cell(row=1, column=col)
        header_cell.fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")  # cor de fundo preta
        header_cell.font = Font(color="FFFFFF", bold=True)  # cor do texto branca e em negrito

messagebox.showinfo("A lista de colaboradores foi gerada com sucesso!", "O arquivo foi salvo na sua pasta de Downloads!")